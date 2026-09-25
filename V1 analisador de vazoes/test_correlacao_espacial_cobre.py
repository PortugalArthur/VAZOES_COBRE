"""Casos de referência para a correlação espacial de lag zero."""

from pathlib import Path
from tempfile import TemporaryDirectory
from io import BytesIO

import duckdb
import numpy as np
import pandas as pd
import pytest
from preparar_ambiente import garantir_dependencias

garantir_dependencias({"openpyxl": "openpyxl>=3.1"})
from openpyxl import load_workbook

from correlacao_espacial_cobre import (
    diagnosticar_par, exportar_correlacao_espacial_xlsx,
    gerar_correlacao_espacial, tabela_espacial,
)


def _arquivo_teste(pasta: Path) -> Path:
    registros = []
    series = {
        0: {0: [1, 2, 3, 4], 1: [4, 3, 2, 1], 2: [1, 2, 1, 2], 3: [7, 7, 7, 7]},
        1: {0: [1, 2, 3, 4], 1: [1, 2, 3, 4], 2: [1, 2, 2, 1], 3: [7, 7, 7, 7]},
    }
    for stage_id, usinas in series.items():
        for hydro_id, valores in usinas.items():
            for scenario_id, valor in enumerate(valores):
                if stage_id == 0 and hydro_id == 1 and scenario_id == 0:
                    valor = np.nan
                registros.append({
                    "scenario_id": scenario_id, "stage_id": stage_id,
                    "hydro_id": hydro_id,
                    "node_id": 1000 + stage_id * 100 + scenario_id * 10 + hydro_id,
                    "incremental_inflow_m3s": valor,
                })
    caminho = pasta / "cenarios.parquet"
    conexao = duckdb.connect()
    try:
        conexao.register("registros", pd.DataFrame(registros).sample(frac=1, random_state=3))
        conexao.execute("COPY registros TO ? (FORMAT PARQUET)", [str(caminho)])
    finally:
        conexao.close()
    return caminho


def test_calculo_tabela_e_exportacao():
    with TemporaryDirectory() as temporario:
        caminho = _arquivo_teste(Path(temporario))
        calendario = {
            0: {"ordem": 0, "mes": 1, "ano": 2026},
            1: {"ordem": 1, "mes": 2, "ano": 2026},
        }
        nomes = {i: f"UHE {i}" for i in range(4)}
        resultado = gerar_correlacao_espacial(str(caminho), nomes, calendario)
        assert resultado.correlacoes.shape == (2, 4, 4)
        assert resultado.correlacoes[0, 0, 1] == pytest.approx(-1)
        assert resultado.correlacoes[1, 0, 1] == pytest.approx(1)
        assert resultado.correlacoes[1, 0, 2] == pytest.approx(0, abs=1e-14)
        assert resultado.correlacoes[0, 0, 2] == pytest.approx(1 / np.sqrt(5))
        assert resultado.correlacoes[0, 1, 0] == resultado.correlacoes[0, 0, 1]
        assert np.isnan(resultado.correlacoes[0, 0, 3])
        assert resultado.pares_validos[0, 0, 1] == 3
        diag = diagnosticar_par(resultado, 0, 0, 1)
        assert diag["Média referência (m³/s)"] == pytest.approx(3)
        assert diag["Média comparada (m³/s)"] == pytest.approx(2)
        assert diag["Pares válidos"] == 3
        assert diag["Status"] == "calculado"
        tabela = tabela_espacial(resultado)
        assert tabela.shape == (8, 7)
        assert np.isnan(tabela.loc[0, "Corr. hydro_id 0"])
        assert tabela.loc[0, "Corr. hydro_id 1"] == pytest.approx(-1)
        somente_ref = gerar_correlacao_espacial(str(caminho), nomes, calendario, 0, 0)
        np.testing.assert_allclose(
            somente_ref.correlacoes[0, 0], resultado.correlacoes[0, 0],
            equal_nan=True, atol=1e-14,
        )
        arquivo = exportar_correlacao_espacial_xlsx(resultado, [0], [0])
        livro = load_workbook(BytesIO(arquivo), read_only=False)
        assert livro["Correlação espacial"].max_row == 2
        assert livro["Valores completos"].sheet_state == "hidden"
        assert livro["Correlação espacial"]["E2"].value == "-1.00"
        assert livro["Correlação espacial"]["D2"].value is None
        assert livro["Valores completos"]["E2"].value == pytest.approx(-1)
        assert livro["Diagnóstico 1"].max_row == 4


def test_pares_insuficientes_e_chave_duplicada():
    with TemporaryDirectory() as temporario:
        pasta = Path(temporario)
        caminho = _arquivo_teste(pasta)
        calendario = {0: {"ordem": 0}, 1: {"ordem": 1}}
        conexao = duckdb.connect()
        try:
            registros = conexao.execute("SELECT * FROM read_parquet(?)", [str(caminho)]).df()
            reduzidos = registros.loc[
                (registros.stage_id.eq(0) & registros.hydro_id.eq(0))
                | (registros.stage_id.eq(0) & registros.hydro_id.eq(1) & registros.scenario_id.eq(1))
            ]
            conexao.register("reduzidos", reduzidos)
            caminho_reduzido = pasta / "reduzido.parquet"
            conexao.execute("COPY reduzidos TO ? (FORMAT PARQUET)", [str(caminho_reduzido)])
            resultado = gerar_correlacao_espacial(str(caminho_reduzido), {}, calendario, 0)
            assert np.isnan(resultado.correlacoes[0, 0, 1])
            assert resultado.estados[0, 0, 1] == 1
            assert resultado.pares_validos[0, 0, 1] == 1
            conexao.register("duplicados", pd.concat([registros, registros.iloc[[0]]], ignore_index=True))
            caminho_duplicado = pasta / "duplicado.parquet"
            conexao.execute("COPY duplicados TO ? (FORMAT PARQUET)", [str(caminho_duplicado)])
        finally:
            conexao.close()
        with pytest.raises(ValueError, match="duplicada"):
            gerar_correlacao_espacial(str(caminho_duplicado), {}, calendario)
