"""Confere ECDF, distância KS, reaproveitamento histórico e XLSX."""

from datetime import date
from io import BytesIO

import numpy as np
import pandas as pd
import pytest

from preparar_ambiente import garantir_dependencias

garantir_dependencias({"openpyxl": "openpyxl>=3.1"})
from openpyxl import load_workbook

from comparacao_ks_cobre import (
    calcular_distancia_ks, exportar_ks_xlsx, gerar_tabela_ks_usina,
)


def test_empates_faltantes_e_local_da_maior_distancia():
    resultado = calcular_distancia_ks([1, 2, 3, 4, np.nan], [1, 1, 1, 1, np.inf])
    assert resultado["cenarios_encontrados"] == 5
    assert resultado["historico_encontrado"] == 5
    assert resultado["cenarios_validos"] == resultado["historico_valido"] == 4
    assert resultado["d"] == pytest.approx(0.75)
    assert resultado["x_max"] == 1
    assert resultado["contagem_cenarios_x"] == 1
    assert resultado["contagem_historico_x"] == 4
    assert resultado["direcao"] == "Maior fração no histórico até x*"

    empate = calcular_distancia_ks([1, 2, 3], [1, 3])
    assert empate["d"] == pytest.approx(1 / 6)
    assert empate["x_max"] == 1  # O menor x vence o empate matemático.
    assert calcular_distancia_ks([-2, -2], [-2])["d"] == 0
    assert calcular_distancia_ks([-2, -2], [-2])["x_max"] is None
    assert calcular_distancia_ks([-2], [3])["d"] == 1
    assert "amostra de apenas um valor" in calcular_distancia_ks([-2], [3])["situacao"]
    assert calcular_distancia_ks([np.nan], [3])["situacao"] == "sem cenários válidos"


def test_tabela_uma_linha_por_estagio_e_historico_nao_repetido():
    calendario = {
        0: {"inicio": date(2026, 1, 1), "mes": 1, "ordem": 0},
        1: {"inicio": date(2026, 2, 1), "mes": 2, "ordem": 1},
        12: {"inicio": date(2027, 1, 1), "mes": 1, "ordem": 2},
    }
    cenarios = pd.DataFrame([
        {"hydro_id": 7, "stage_id": stage_id, "scenario_id": scenario_id,
         "incremental_inflow_m3s": valor}
        for stage_id, valores in ((0, [1, 2, 3, 4]), (1, [1, 2, 3, 4]), (12, [1, 1, 1, 1]))
        for scenario_id, valor in enumerate(valores)
    ])
    historico = pd.DataFrame({
        "hydro_id": [7, 7],
        "start_date": [date(2020, 1, 1), date(2021, 1, 1)],
        "value_m3s": [1, 1],
    })
    tabela = gerar_tabela_ks_usina(cenarios, historico, calendario, 7, "Teste")
    assert tabela.stage_id.tolist() == [0, 1, 12]
    assert tabela.historico_encontrado.tolist() == [2, 0, 2]
    assert tabela.d.iloc[0] == pytest.approx(0.75)
    assert pd.isna(tabela.d.iloc[1])
    assert tabela.situacao.iloc[1] == "sem histórico válido para o mês"
    assert tabela.d.iloc[2] == 0
    assert pd.isna(tabela.x_max.iloc[2])

    arquivo = load_workbook(BytesIO(exportar_ks_xlsx(tabela)))
    assert arquivo["Valores numéricos"].sheet_state == "hidden"
    folha = arquivo["Comparação KS"]
    assert folha.max_row == 4
    assert folha["I2"].value == "0.750000"
    assert folha["I3"].value is None
    assert arquivo["Valores numéricos"]["I2"].value == pytest.approx(0.75)
    assert folha["A1"].fill.fgColor.rgb.endswith("09B6CB")


def test_chave_duplicada_e_mes_historico_duplicado_sao_rejeitados():
    calendario = {0: {"inicio": date(2026, 1, 1), "mes": 1, "ordem": 0}}
    cenarios = pd.DataFrame({
        "hydro_id": [7, 7], "stage_id": [0, 0], "scenario_id": [1, 1],
        "incremental_inflow_m3s": [1, 2],
    })
    historico = pd.DataFrame({
        "hydro_id": [7], "start_date": [date(2020, 1, 1)], "value_m3s": [1],
    })
    with pytest.raises(ValueError, match="duplicados"):
        gerar_tabela_ks_usina(cenarios, historico, calendario, 7, "Teste")
    with pytest.raises(ValueError, match="duplicados"):
        gerar_tabela_ks_usina(cenarios.iloc[[0]], pd.concat([historico, historico]), calendario, 7, "Teste")
