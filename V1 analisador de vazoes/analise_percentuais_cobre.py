"""Fatias acumuladas dos menores valores para gráficos e relatório COBRE."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
import math

import numpy as np
import pandas as pd

from relatorio_estatistico_cobre import JANELAS, MESES


PERCENTUAIS = tuple(range(5, 100, 5))
VERSAO_ESQUEMA_PERCENTUAIS = 5

COLUNAS_PERCENTUAIS = [
    ("usina", "UHE"),
    ("hydro_id", "hydro_id"),
    ("fonte", "Fonte"),
    ("janela", "Janela da análise"),
    ("mes_nome", "Mês histórico"),
    ("stage_id", "stage_id"),
    ("percentual", "Percentual dos menores valores"),
    ("n_validos", "Observações válidas no grupo"),
    ("n_selecionados", "Observações selecionadas"),
    ("media", "Média (m³/s)"),
    ("desvio", "Desvio padrão populacional (m³/s)"),
    ("assimetria", "Assimetria"),
    ("curtose", "Curtose de Fisher"),
    ("minimo", "Mínimo (m³/s)"),
    ("q1", "Q1 (m³/s)"),
    ("mediana", "Mediana (m³/s)"),
    ("q3", "Q3 (m³/s)"),
    ("maximo", "Máximo (m³/s)"),
    ("haste_inferior", "Haste inferior (m³/s)"),
    ("haste_superior", "Haste superior (m³/s)"),
    ("quantidade_outliers", "Quantidade de outliers"),
]

COLUNAS_INTEIRAS = {
    "hydro_id", "stage_id", "n_validos", "n_selecionados", "quantidade_outliers",
}
COLUNAS_DECIMAIS = {
    "media", "desvio", "assimetria", "curtose", "minimo", "q1", "mediana", "q3", "maximo",
    "haste_inferior", "haste_superior",
}


def _quartil_plotly_linear(ordenados: np.ndarray, fracao: float) -> float:
    """Reproduz Plotly.js Lib.interp: posição fracao * N - 0,5."""
    # https://github.com/plotly/plotly.js/blob/master/src/lib/stats.js
    posicao = fracao * len(ordenados) - 0.5
    if posicao <= 0:
        return float(ordenados[0])
    if posicao >= len(ordenados) - 1:
        return float(ordenados[-1])
    inferior = math.floor(posicao)
    peso = posicao - inferior
    return float((1 - peso) * ordenados[inferior] + peso * ordenados[inferior + 1])


def _fatiar_ordenados(ordenados: np.ndarray, percentual: int) -> np.ndarray:
    """Retorna a mesma fatia inferior acumulada usada em todos os cálculos."""
    k = (len(ordenados) * percentual + 99) // 100
    return ordenados[:k]


def _resumir_fatia(ordenados: np.ndarray, percentual: int) -> dict:
    """Calcula todos os números exibidos a partir da mesma fatia inferior."""
    n = int(len(ordenados))
    fatia = _fatiar_ordenados(ordenados, percentual)
    k = int(len(fatia))
    resultado = {
        "percentual": percentual,
        "n_validos": n,
        "n_selecionados": k,
        "media": None,
        "desvio": None,
        "assimetria": None,
        "curtose": None,
        "minimo": None,
        "q1": None,
        "mediana": None,
        "q3": None,
        "maximo": None,
        "haste_inferior": None,
        "haste_superior": None,
        "quantidade_outliers": None,
        "outliers": (),
    }
    if not k:
        return resultado

    media = float(np.mean(fatia))
    desvios = fatia - media
    m2 = float(np.mean(desvios ** 2))
    desvio = math.sqrt(m2)
    if k >= 2 and math.isfinite(desvio) and desvio > 0:
        with np.errstate(over="ignore", invalid="ignore", divide="ignore"):
            padronizados = desvios / desvio
            assimetria = float(np.mean(padronizados ** 3))
            curtose = float(np.mean(padronizados ** 4) - 3)
        resultado["assimetria"] = assimetria if math.isfinite(assimetria) else None
        resultado["curtose"] = curtose if math.isfinite(curtose) else None
    q1 = _quartil_plotly_linear(fatia, 0.25)
    mediana = _quartil_plotly_linear(fatia, 0.50)
    q3 = _quartil_plotly_linear(fatia, 0.75)
    amplitude = q3 - q1
    limite_inferior = q1 - 1.5 * amplitude
    limite_superior = q3 + 1.5 * amplitude
    dentro = fatia[(fatia >= limite_inferior) & (fatia <= limite_superior)]
    # A pequena tolerância usada pelo Plotly no limite não muda os quartis;
    # estes valores de haste são fornecidos ao gráfico já calculados.
    haste_inferior = float(dentro[0]) if dentro.size else q1
    haste_superior = float(dentro[-1]) if dentro.size else q3
    outliers = tuple(float(valor) for valor in fatia[
        (fatia < haste_inferior) | (fatia > haste_superior)
    ])
    resultado.update(
        media=media,
        desvio=desvio,
        minimo=float(fatia[0]),
        q1=q1,
        mediana=mediana,
        q3=q3,
        maximo=float(fatia[-1]),
        haste_inferior=haste_inferior,
        haste_superior=haste_superior,
        quantidade_outliers=len(outliers),
        outliers=outliers,
    )
    return resultado


def _ordenar_finitos(valores: pd.Series, desempate: pd.Series | None = None) -> np.ndarray:
    """Ordena valores finitos e usa uma chave estável quando a identidade importa."""
    tabela = pd.DataFrame({"valor": pd.to_numeric(valores, errors="coerce")})
    tabela = tabela.loc[np.isfinite(tabela["valor"])].copy()
    if desempate is not None:
        tabela["desempate"] = desempate.reindex(tabela.index)
        tabela = tabela.sort_values(["valor", "desempate"], kind="stable", na_position="last")
    else:
        tabela = tabela.sort_values("valor", kind="stable")
    return tabela["valor"].to_numpy(dtype=float)


def selecionar_menores_valores(valores: pd.Series, percentual: int) -> np.ndarray:
    """Entrega os valores exatos da fatia inferior usada no resumo percentual."""
    if percentual not in PERCENTUAIS:
        raise ValueError(f"Percentual inválido: p{percentual}%.")
    return _fatiar_ordenados(_ordenar_finitos(valores), percentual).copy()


def gerar_resumo_percentuais(
    cenarios: pd.DataFrame,
    historico: pd.DataFrame | None,
    hydro_id: int,
    nome_usina: str,
    cenarios_somados: dict[int, pd.DataFrame] | None = None,
    historicos_somados: dict[int, pd.DataFrame] | None = None,
    estagios_ordenados: list[int] | None = None,
) -> pd.DataFrame:
    """Cria as fatias das vazões originais e das somas móveis de uma UHE."""
    linhas = []
    cenarios_por_janela = {JANELAS[0]: (cenarios, "incremental_inflow_m3s")}
    historicos_por_janela = (
        {JANELAS[0]: historico} if historico is not None else {}
    )
    for indice, tamanho in enumerate((6, 12), start=1):
        if cenarios_somados and tamanho in cenarios_somados:
            cenarios_por_janela[JANELAS[indice]] = (
                cenarios_somados[tamanho], "soma_vazoes_m3s"
            )
        if historicos_somados and tamanho in historicos_somados:
            historicos_por_janela[JANELAS[indice]] = historicos_somados[tamanho]

    vazio = np.array([], dtype=float)
    for janela in JANELAS:
        dados_historicos = historicos_por_janela.get(janela)
        if dados_historicos is None:
            continue
        grupos_historicos = {}
        for mes, grupo in dados_historicos.groupby("mes", sort=False):
            if pd.isna(mes):
                continue
            desempate = grupo["start_date"] if "start_date" in grupo.columns else None
            grupos_historicos[int(mes)] = _ordenar_finitos(grupo["value_m3s"], desempate)
        for mes in range(1, 13):
            ordenados = grupos_historicos.get(mes, vazio)
            for percentual in PERCENTUAIS:
                linhas.append({
                    "usina": nome_usina,
                    "hydro_id": hydro_id,
                    "fonte": "Histórico",
                    "janela": janela,
                    "mes_num": mes,
                    "mes_nome": MESES[mes],
                    "stage_id": None,
                    **_resumir_fatia(ordenados, percentual),
                })

    for janela in JANELAS:
        analise = cenarios_por_janela.get(janela)
        if analise is None:
            continue
        dados_cenarios, coluna_valor = analise
        grupos = {int(chave): grupo for chave, grupo in dados_cenarios.groupby("stage_id", sort=False)}
        ordem = (
            [stage_id for stage_id in estagios_ordenados if stage_id in grupos]
            if estagios_ordenados is not None else sorted(grupos)
        )
        for stage_id in ordem:
            grupo = grupos[stage_id]
            desempate = grupo["scenario_id"] if "scenario_id" in grupo.columns else None
            ordenados = _ordenar_finitos(grupo[coluna_valor], desempate)
            for percentual in PERCENTUAIS:
                linhas.append({
                    "usina": nome_usina,
                    "hydro_id": hydro_id,
                    "fonte": "Cenários",
                    "janela": janela,
                    "mes_num": None,
                    "mes_nome": "",
                    "stage_id": int(stage_id),
                    **_resumir_fatia(ordenados, percentual),
                })
    return pd.DataFrame(linhas)


def filtrar_percentuais(
    resumo: pd.DataFrame,
    percentuais: list[int],
    janelas: list[str] | None = None,
) -> pd.DataFrame:
    """Filtra percentuais e janelas sem recalcular o resumo."""
    mascara = resumo["percentual"].isin(percentuais)
    if janelas:
        mascara &= resumo["janela"].isin(janelas)
    return resumo.loc[mascara].reset_index(drop=True)


def tabela_percentuais_visivel(resumo: pd.DataFrame) -> pd.DataFrame:
    """Prepara nomes legíveis e mantém valores numéricos na tabela Streamlit."""
    visivel = resumo[[chave for chave, _ in COLUNAS_PERCENTUAIS]].copy()
    visivel.columns = [nome for _, nome in COLUNAS_PERCENTUAIS]
    visivel["stage_id"] = visivel["stage_id"].astype("Int64")
    visivel["Observações válidas no grupo"] = visivel["Observações válidas no grupo"].astype("Int64")
    visivel["Observações selecionadas"] = visivel["Observações selecionadas"].astype("Int64")
    visivel["Quantidade de outliers"] = visivel["Quantidade de outliers"].astype("Int64")
    visivel["Percentual dos menores valores"] = visivel["Percentual dos menores valores"].map(
        lambda percentual: f"p{int(percentual)}%"
    )
    return visivel


def _numero_visivel(valor) -> str | None:
    if valor is None or pd.isna(valor):
        return None
    numero = float(valor)
    if not math.isfinite(numero):
        return None
    texto = f"{numero:.2f}".rstrip("0").rstrip(".")
    return "0" if texto in ("", "-0") else texto


def _valor_xlsx(chave: str, valor, numerico: bool):
    if valor is None or pd.isna(valor):
        return None
    if chave == "percentual":
        return int(valor) if numerico else f"p{int(valor)}%"
    if chave in COLUNAS_INTEIRAS:
        return int(valor)
    if chave in COLUNAS_DECIMAIS:
        return float(valor) if numerico else _numero_visivel(valor)
    return valor


def exportar_percentuais_xlsx(resumo: pd.DataFrame) -> bytes:
    """Exporta uma UHE e todos os percentuais com o tema da primeira tabela."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    planilha = livro.active
    planilha.title = "Fatias inferiores"
    numerica = livro.create_sheet("Valores numéricos")
    cabecalhos = [nome for _, nome in COLUNAS_PERCENTUAIS]
    ciano = "09B6CB"
    ciano_escuro = "067F94"
    fina = Side(style="thin", color="B3DEE5")
    grossa = Side(style="medium", color=ciano_escuro)
    borda = Border(left=fina, right=fina, top=fina, bottom=fina)
    separador = Border(left=fina, right=fina, top=grossa, bottom=fina)
    cabecalho = PatternFill("solid", fgColor=ciano)
    historico = PatternFill("solid", fgColor="ECF9FB")
    cenarios = PatternFill("solid", fgColor="FFFFFF")
    fonte_cabecalho = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    fonte_corpo = Font(name="Aptos", size=10, color="183444")
    texto = Alignment(horizontal="left", vertical="center")
    numero = Alignment(horizontal="right", vertical="center")
    larguras = [len(nome) + 2 for nome in cabecalhos]

    for folha in (planilha, numerica):
        folha.append(cabecalhos)
        folha.sheet_view.showGridLines = False
        folha.freeze_panes = "C2"
    for celula in planilha[1]:
        celula.fill = cabecalho
        celula.font = fonte_cabecalho
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celula.border = borda
    planilha.row_dimensions[1].height = 43

    anterior = None
    for indice, registro in enumerate(resumo.itertuples(index=False), start=2):
        dados = registro._asdict()
        planilha.append([_valor_xlsx(chave, dados[chave], False) for chave, _ in COLUNAS_PERCENTUAIS])
        numerica.append([_valor_xlsx(chave, dados[chave], True) for chave, _ in COLUNAS_PERCENTUAIS])
        troca_fonte = anterior == "Histórico" and dados["fonte"] == "Cenários"
        preenchimento = historico if dados["fonte"] == "Histórico" else cenarios
        for posicao, (chave, _) in enumerate(COLUNAS_PERCENTUAIS, start=1):
            celula = planilha.cell(row=indice, column=posicao)
            celula.font = fonte_corpo
            celula.fill = preenchimento
            celula.alignment = numero if chave in COLUNAS_INTEIRAS | COLUNAS_DECIMAIS else texto
            celula.border = separador if troca_fonte else borda
            larguras[posicao - 1] = max(larguras[posicao - 1], len(str(celula.value or "")) + 2)
        planilha.row_dimensions[indice].height = 19
        anterior = dados["fonte"]

    for posicao, largura in enumerate(larguras, start=1):
        planilha.column_dimensions[get_column_letter(posicao)].width = min(max(largura, 12), 80)
    ultima_coluna = get_column_letter(len(COLUNAS_PERCENTUAIS))
    planilha.auto_filter.ref = f"A1:{ultima_coluna}{max(1, planilha.max_row)}"
    planilha.sheet_properties.tabColor = ciano_escuro
    numerica.sheet_state = "hidden"
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def nome_arquivo_percentuais(hydro_id: int) -> str:
    return f"relatorio_percentuais_cobre_uhe_{hydro_id}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
