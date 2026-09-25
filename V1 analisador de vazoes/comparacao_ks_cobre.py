"""Comparação descritiva de Kolmogorov–Smirnov entre cenários e histórico mensal."""

from datetime import datetime
from io import BytesIO

import numpy as np
import pandas as pd


VERSAO_COMPARACAO_KS = 1

COLUNAS_KS = [
    ("usina", "UHE"),
    ("hydro_id", "hydro_id"),
    ("stage_id", "stage_id"),
    ("mes_ano", "Mês/ano"),
    ("cenarios_encontrados", "Cenários encontrados"),
    ("cenarios_validos", "Cenários válidos"),
    ("historico_encontrado", "Histórico encontrado no mês"),
    ("historico_valido", "Histórico válido no mês"),
    ("d", "Distância KS (D)"),
    ("d_pp", "Diferença máxima (p.p.)"),
    ("x_max", "Vazão da diferença máxima (m³/s)"),
    ("contagem_cenarios_x", "Cenários acumulados em x*"),
    ("contagem_historico_x", "Histórico acumulado em x*"),
    ("acumulado_cenarios_pct", "Cenários acumulados em x* (%)"),
    ("acumulado_historico_pct", "Histórico acumulado em x* (%)"),
    ("direcao", "Direção da diferença"),
    ("situacao", "Situação"),
]

COLUNAS_INTEIRAS = {
    "hydro_id", "stage_id", "cenarios_encontrados", "cenarios_validos",
    "historico_encontrado", "historico_valido", "contagem_cenarios_x",
    "contagem_historico_x",
}
COLUNAS_DECIMAIS = {
    "d", "d_pp", "x_max", "acumulado_cenarios_pct", "acumulado_historico_pct",
}


def _numeros_finitos(valores) -> tuple[np.ndarray, int]:
    serie = pd.to_numeric(pd.Series(valores), errors="coerce")
    encontrados = len(serie)
    numeros = serie.to_numpy(dtype=float, na_value=np.nan)
    return np.sort(numeros[np.isfinite(numeros)]), encontrados


def calcular_distancia_ks(cenarios, historico) -> dict:
    """Calcula ECDFs e maior distância no suporte observado, sem KDE ou valor-p."""
    c, encontrados_c = _numeros_finitos(cenarios)
    h, encontrados_h = _numeros_finitos(historico)
    resultado = {
        "cenarios_encontrados": encontrados_c,
        "cenarios_validos": len(c),
        "historico_encontrado": encontrados_h,
        "historico_valido": len(h),
        "valores_x": np.array([], dtype=float),
        "contagens_cenarios": np.array([], dtype=int),
        "contagens_historico": np.array([], dtype=int),
        "acumulados_cenarios": np.array([], dtype=float),
        "acumulados_historico": np.array([], dtype=float),
        "d": None,
        "x_max": None,
        "contagem_cenarios_x": None,
        "contagem_historico_x": None,
        "acumulado_cenarios_x": None,
        "acumulado_historico_x": None,
        "direcao": None,
        "situacao": "calculado",
    }
    if len(c) == 0:
        resultado["situacao"] = "sem cenários válidos"
        return resultado
    if len(h) == 0:
        resultado["situacao"] = "sem histórico válido para o mês"
        return resultado
    if min(len(c), len(h)) == 1:
        resultado["situacao"] = "calculado com amostra de apenas um valor"

    x = np.union1d(c, h)
    contagens_c = np.searchsorted(c, x, side="right")
    contagens_h = np.searchsorted(h, x, side="right")
    acumulados_c = contagens_c / len(c)
    acumulados_h = contagens_h / len(h)
    diferencas = acumulados_c - acumulados_h
    distancias = np.abs(diferencas)
    maior = float(np.max(distancias))
    indice = int(np.flatnonzero(np.isclose(distancias, maior, rtol=0, atol=4 * np.finfo(float).eps))[0])
    d = float(abs(diferencas[indice]))
    resultado.update({
        "valores_x": x,
        "contagens_cenarios": contagens_c,
        "contagens_historico": contagens_h,
        "acumulados_cenarios": acumulados_c,
        "acumulados_historico": acumulados_h,
        "d": d,
    })
    if d > 0:
        resultado.update({
            "x_max": float(x[indice]),
            "contagem_cenarios_x": int(contagens_c[indice]),
            "contagem_historico_x": int(contagens_h[indice]),
            "acumulado_cenarios_x": float(acumulados_c[indice]),
            "acumulado_historico_x": float(acumulados_h[indice]),
            "direcao": (
                "Maior fração nos cenários até x*"
                if diferencas[indice] > 0 else "Maior fração no histórico até x*"
            ),
        })
    else:
        resultado["direcao"] = "Curvas acumuladas coincidentes"
    return resultado


def gerar_tabela_ks_usina(
    cenarios: pd.DataFrame, historico: pd.DataFrame,
    calendario: dict[int, dict], hydro_id: int, nome_usina: str,
) -> pd.DataFrame:
    """Produz uma linha por estágio, usando cada mês histórico original uma vez."""
    if not cenarios.empty and not cenarios["hydro_id"].eq(hydro_id).all():
        raise ValueError("Os cenários da tabela KS contêm outra UHE.")
    if cenarios.duplicated(["hydro_id", "stage_id", "scenario_id"]).any():
        raise ValueError("Há cenários duplicados para UHE, estágio e scenario_id.")
    if not historico.empty and not historico["hydro_id"].eq(hydro_id).all():
        raise ValueError("O histórico da tabela KS contém outra UHE.")

    datas = pd.to_datetime(historico["start_date"], errors="coerce")
    if datas.isna().any():
        raise ValueError("O histórico possui mês/ano inválido para a tabela KS.")
    periodos = datas.dt.to_period("M")
    if periodos.duplicated().any():
        raise ValueError("O histórico possui UHE e mês/ano duplicados para a tabela KS.")
    historico_por_mes = {
        mes: historico.loc[datas.dt.month.eq(mes), "value_m3s"]
        for mes in range(1, 13)
    }
    stage_ids = [int(valor) for valor in cenarios["stage_id"].dropna().unique()]
    ausentes = sorted(set(stage_ids) - set(calendario))
    if ausentes:
        raise ValueError(f"O stages.json não contém os estágios da tabela KS: {ausentes}.")

    linhas = []
    for stage_id in sorted(stage_ids, key=lambda valor: calendario[valor]["ordem"]):
        registro = calendario[stage_id]
        comparacao = calcular_distancia_ks(
            cenarios.loc[cenarios["stage_id"].eq(stage_id), "incremental_inflow_m3s"],
            historico_por_mes[registro["mes"]],
        )
        d = comparacao["d"]
        linhas.append({
            "usina": nome_usina,
            "hydro_id": hydro_id,
            "stage_id": stage_id,
            "mes_ano": f"{registro['inicio'].month:02d}/{registro['inicio'].year}",
            "cenarios_encontrados": comparacao["cenarios_encontrados"],
            "cenarios_validos": comparacao["cenarios_validos"],
            "historico_encontrado": comparacao["historico_encontrado"],
            "historico_valido": comparacao["historico_valido"],
            "d": d,
            "d_pp": None if d is None else 100 * d,
            "x_max": comparacao["x_max"],
            "contagem_cenarios_x": comparacao["contagem_cenarios_x"],
            "contagem_historico_x": comparacao["contagem_historico_x"],
            "acumulado_cenarios_pct": (
                None if comparacao["acumulado_cenarios_x"] is None
                else 100 * comparacao["acumulado_cenarios_x"]
            ),
            "acumulado_historico_pct": (
                None if comparacao["acumulado_historico_x"] is None
                else 100 * comparacao["acumulado_historico_x"]
            ),
            "direcao": comparacao["direcao"],
            "situacao": comparacao["situacao"],
        })
    return pd.DataFrame(linhas, columns=[chave for chave, _ in COLUNAS_KS])


def tabela_ks_visivel(tabela: pd.DataFrame) -> pd.DataFrame:
    visivel = tabela[[chave for chave, _ in COLUNAS_KS]].copy()
    visivel.columns = [nome for _, nome in COLUNAS_KS]
    return visivel


def _valor_excel(chave: str, valor, numerico: bool):
    if pd.isna(valor):
        return None
    if chave in COLUNAS_INTEIRAS:
        return int(valor)
    if chave in COLUNAS_DECIMAIS:
        numero = float(valor)
        if numerico:
            return numero
        casas = 6 if chave == "d" else 2
        texto = f"{numero:.{casas}f}"
        return "0" if float(texto) == 0 else texto
    return str(valor)


def exportar_ks_xlsx(tabela: pd.DataFrame) -> bytes:
    """Exporta a UHE completa em ciano e preserva os números em folha oculta."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    folha = livro.active
    folha.title = "Comparação KS"
    numerica = livro.create_sheet("Valores numéricos")
    cabecalhos = [nome for _, nome in COLUNAS_KS]
    fina = Side(style="thin", color="B3DEE5")
    borda = Border(left=fina, right=fina, top=fina, bottom=fina)
    ciano = PatternFill("solid", fgColor="09B6CB")
    listrado = PatternFill("solid", fgColor="ECF9FB")
    branco = PatternFill("solid", fgColor="FFFFFF")
    fonte_cabecalho = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    fonte_corpo = Font(name="Aptos", size=10, color="183444")
    larguras = [len(nome) + 2 for nome in cabecalhos]
    for planilha in (folha, numerica):
        planilha.append(cabecalhos)
        planilha.sheet_view.showGridLines = False
        planilha.freeze_panes = "E2"
    for celula in folha[1]:
        celula.fill = ciano
        celula.font = fonte_cabecalho
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celula.border = borda
    folha.row_dimensions[1].height = 45
    for numero_linha, (_, registro) in enumerate(tabela.iterrows(), start=2):
        folha.append([_valor_excel(chave, registro[chave], False) for chave, _ in COLUNAS_KS])
        numerica.append([_valor_excel(chave, registro[chave], True) for chave, _ in COLUNAS_KS])
        for coluna, (chave, _) in enumerate(COLUNAS_KS, start=1):
            celula = folha.cell(numero_linha, coluna)
            celula.fill = listrado if numero_linha % 2 == 0 else branco
            celula.font = fonte_corpo
            celula.border = borda
            celula.alignment = Alignment(
                horizontal="right" if chave in COLUNAS_INTEIRAS | COLUNAS_DECIMAIS else "left",
                vertical="center",
            )
            larguras[coluna - 1] = max(larguras[coluna - 1], len(str(celula.value or "")) + 2)
        folha.row_dimensions[numero_linha].height = 20
    for coluna, largura in enumerate(larguras, start=1):
        folha.column_dimensions[get_column_letter(coluna)].width = min(max(largura, 12), 65)
    ultima = get_column_letter(len(COLUNAS_KS))
    folha.auto_filter.ref = f"A1:{ultima}{max(1, folha.max_row)}"
    folha.sheet_properties.tabColor = "067F94"
    numerica.sheet_state = "hidden"
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def nome_arquivo_ks(hydro_id: int) -> str:
    return f"comparacao_ks_cobre_uhe_{hydro_id}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
