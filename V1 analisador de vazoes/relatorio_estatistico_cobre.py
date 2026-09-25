"""Resumo estatístico auditável para a segunda aba do analisador COBRE."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
import json
import math

import duckdb
import numpy as np
import pandas as pd


MESES = {
    1: "Janeiro", 2: "Fevereiro", 3: "Março", 4: "Abril",
    5: "Maio", 6: "Junho", 7: "Julho", 8: "Agosto",
    9: "Setembro", 10: "Outubro", 11: "Novembro", 12: "Dezembro",
}

COLUNAS = [
    ("usina", "UHE"),
    ("hydro_id", "hydro_id"),
    ("fonte", "Fonte"),
    ("janela", "Janela da análise"),
    ("mes_nome", "Mês histórico"),
    ("stage_id", "stage_id"),
    ("n_validos", "Observações válidas"),
    ("media", "Média (m³/s)"),
    ("desvio", "Desvio padrão populacional (m³/s)"),
    ("assimetria", "Assimetria"),
    ("curtose", "Curtose de Fisher"),
    ("minimo", "Mínimo (m³/s)"),
    ("maximo", "Máximo (m³/s)"),
]
COLUNAS_NUMERICAS = {"media", "desvio", "assimetria", "curtose", "minimo", "maximo"}
JANELAS = ("Sem soma", "6 estágios/meses", "12 estágios/meses")


def _numero_finito(valor) -> float | None:
    """Converte resultados agregados nulos ou não finitos em ausência explícita."""
    if valor is None or pd.isna(valor):
        return None
    numero = float(valor)
    return numero if math.isfinite(numero) else None


def _proximo_mes(data: date) -> date:
    """Retorna o primeiro dia do mês seguinte."""
    return date(data.year + (data.month == 12), 1 if data.month == 12 else data.month + 1, 1)


def _carregar_calendario(caminho: str) -> pd.DataFrame:
    """Lê e valida a ordem mensal usada nas janelas de cenários."""
    try:
        with open(caminho, "r", encoding="utf-8") as arquivo:
            documento = json.load(arquivo)
    except (json.JSONDecodeError, UnicodeError) as erro:
        raise ValueError("O stages.json não contém um JSON válido em UTF-8.") from erro
    if not isinstance(documento, dict) or not isinstance(documento.get("stages"), list):
        raise ValueError("Estrutura inválida: o stages.json deve conter uma lista 'stages'.")
    registros = []
    ids = set()
    for posicao, estagio in enumerate(documento["stages"], start=1):
        if not isinstance(estagio, dict) or any(campo not in estagio for campo in ("id", "start_date", "end_date")):
            raise ValueError(f"Estágio {posicao} do stages.json não possui id, start_date e end_date válidos.")
        stage_id = estagio["id"]
        if isinstance(stage_id, bool) or not isinstance(stage_id, int) or stage_id in ids:
            raise ValueError(f"stage_id inválido ou repetido no stages.json: {stage_id!r}.")
        try:
            inicio = date.fromisoformat(estagio["start_date"])
            fim = date.fromisoformat(estagio["end_date"])
        except (TypeError, ValueError) as erro:
            raise ValueError(f"stage_id {stage_id} possui datas inválidas no stages.json.") from erro
        if inicio.day != 1 or fim != _proximo_mes(inicio):
            raise ValueError(f"stage_id {stage_id} não representa um mês calendário completo.")
        ids.add(stage_id)
        registros.append({"stage_id": stage_id, "start_date": inicio, "end_date": fim})
    if not registros:
        raise ValueError("O stages.json não possui estágios.")
    registros.sort(key=lambda item: item["start_date"])
    for anterior, atual in zip(registros, registros[1:]):
        if atual["start_date"] != anterior["end_date"]:
            raise ValueError(
                f"Há lacuna ou sobreposição entre os stage_id {anterior['stage_id']} e {atual['stage_id']}."
            )
    for ordem, registro in enumerate(registros):
        registro["ordem"] = ordem
        registro["mes_num"] = registro["start_date"].month
    return pd.DataFrame(registros)


def _resumo_cenarios_janela(
    conexao: duckdb.DuckDBPyConnection,
    caminho: str,
    janela: int,
) -> pd.DataFrame:
    """Calcula as somas completas do mesmo cenário e as resume no estágio terminal."""
    return conexao.execute(
        f"""
        WITH dados AS (
            SELECT d.hydro_id, d.scenario_id, d.stage_id, c.ordem,
                   CASE WHEN isfinite(d.incremental_inflow_m3s)
                        THEN d.incremental_inflow_m3s END AS vazao
            FROM read_parquet(?) AS d
            JOIN calendario AS c USING (stage_id)
        ),
        janelas AS (
            SELECT *,
                   SUM(vazao) OVER janela AS soma,
                   COUNT(vazao) OVER janela AS valores_validos,
                   COUNT(*) OVER janela AS registros_janela,
                   MIN(ordem) OVER janela AS primeira_ordem
            FROM dados
            WINDOW janela AS (
                PARTITION BY hydro_id, scenario_id
                ORDER BY ordem
                ROWS BETWEEN {janela - 1} PRECEDING AND CURRENT ROW
            )
        ),
        validas AS (
            SELECT hydro_id, stage_id, soma
            FROM janelas
            WHERE ordem >= {janela - 1}
              AND valores_validos = {janela}
              AND registros_janela = {janela}
              AND ordem - primeira_ordem = {janela - 1}
        ),
        centradas AS (
            SELECT *, AVG(soma) OVER (PARTITION BY hydro_id, stage_id) AS media_grupo
            FROM validas
        )
        SELECT hydro_id, stage_id,
               COUNT(*) AS n_validos,
               AVG(soma) AS media,
               STDDEV_POP(soma) AS desvio,
               AVG(POWER(soma - media_grupo, 3)) AS m3,
               AVG(POWER(soma - media_grupo, 4)) AS m4,
               MIN(soma) AS minimo,
               MAX(soma) AS maximo
        FROM centradas
        GROUP BY hydro_id, stage_id
        ORDER BY hydro_id, stage_id
        """,
        [caminho],
    ).df()


def _resumo_cenarios(caminho: str, calendario: pd.DataFrame) -> pd.DataFrame:
    """Agrega os valores originais e as duas janelas sem carregar o Parquet no Pandas."""
    conexao = duckdb.connect(database=":memory:")
    try:
        conexao.register("calendario", calendario[["stage_id", "ordem"]])
        base = conexao.execute(
            """
            WITH dados AS (
                SELECT hydro_id, stage_id, scenario_id, node_id,
                       CASE WHEN isfinite(incremental_inflow_m3s)
                            THEN incremental_inflow_m3s END AS vazao
                FROM read_parquet(?)
            ),
            centrados AS (
                SELECT *, AVG(vazao) OVER (PARTITION BY hydro_id, stage_id) AS media_grupo
                FROM dados
            )
            SELECT hydro_id, stage_id,
                   COUNT(*) AS registros,
                   COUNT(DISTINCT scenario_id) AS cenarios_distintos,
                   COUNT(node_id) AS nos_identificados,
                   COUNT(vazao) AS n_validos,
                   AVG(vazao) AS media,
                   STDDEV_POP(vazao) AS desvio,
                   AVG(POWER(vazao - media_grupo, 3)) AS m3,
                   AVG(POWER(vazao - media_grupo, 4)) AS m4,
                   MIN(vazao) AS minimo,
                   MAX(vazao) AS maximo
            FROM centrados
            GROUP BY hydro_id, stage_id
            ORDER BY hydro_id, stage_id
            """,
            [caminho],
        ).df()
        if base.empty:
            raise ValueError("O Parquet consolidado não possui UHEs e estágios para o relatório.")
        if base[["hydro_id", "stage_id"]].isna().any().any():
            raise ValueError("O Parquet consolidado possui hydro_id ou stage_id nulo.")
        invalidos = base.loc[
            (base["registros"] != base["cenarios_distintos"])
            | (base["registros"] != base["nos_identificados"])
        ]
        if not invalidos.empty:
            exemplo = invalidos.iloc[0]
            raise ValueError(
                f"Cenários repetidos ou identificadores ausentes em hydro_id "
                f"{int(exemplo['hydro_id'])}, stage_id {int(exemplo['stage_id'])}."
            )
        sem_calendario = sorted(set(base["stage_id"].astype(int)) - set(calendario["stage_id"].astype(int)))
        if sem_calendario:
            amostra = ", ".join(str(valor) for valor in sem_calendario[:20])
            raise ValueError(f"O stages.json não possui os stage_id do Parquet: {amostra}.")
        base["janela"] = JANELAS[0]
        base["janela_ordem"] = 0
        base = base.drop(columns=["registros", "cenarios_distintos", "nos_identificados"])
        partes = [base]
        for ordem, tamanho_janela in enumerate((6, 12), start=1):
            resumo = _resumo_cenarios_janela(conexao, caminho, tamanho_janela)
            resumo["janela"] = JANELAS[ordem]
            resumo["janela_ordem"] = ordem
            partes.append(resumo)
        return pd.concat(partes, ignore_index=True)
    finally:
        conexao.close()


def _resumir_valores(valores) -> dict:
    """Aplica os mesmos momentos populacionais a qualquer conjunto da tabela."""
    numericos = pd.to_numeric(pd.Series(valores), errors="coerce").to_numpy(dtype=float)
    numericos = numericos[np.isfinite(numericos)]
    n = int(numericos.size)
    resumo = {
        "n_validos": n, "media": None, "desvio": None,
        "assimetria": None, "curtose": None, "minimo": None, "maximo": None,
    }
    if not n:
        return resumo
    media = float(np.mean(numericos))
    desvios = numericos - media
    m2 = float(np.mean(desvios ** 2))
    resumo.update(
        media=media,
        desvio=math.sqrt(m2),
        minimo=float(np.min(numericos)),
        maximo=float(np.max(numericos)),
    )
    if n >= 2 and m2 > 0:
        sigma = math.sqrt(m2)
        resumo["assimetria"] = float(np.mean(desvios ** 3) / sigma ** 3)
        resumo["curtose"] = float(np.mean(desvios ** 4) / sigma ** 4 - 3)
    return resumo


def _resumo_historico(caminho: str) -> dict[tuple[int, str, int], dict]:
    """Resume valores mensais e janelas móveis históricas pela data terminal."""
    conexao = duckdb.connect(database=":memory:")
    try:
        dados = conexao.execute(
            """
            SELECT hydro_id,
                   TRY_CAST(start_date AS DATE) AS start_date,
                   TRY_CAST(value_m3s AS DOUBLE) AS vazao
            FROM read_parquet(?)
            ORDER BY hydro_id, start_date
            """,
            [caminho],
        ).df()
    finally:
        conexao.close()
    if dados[["hydro_id", "start_date"]].isna().any().any():
        raise ValueError("O histórico possui hydro_id ou start_date ausente ou inválido.")
    dados["periodo"] = pd.to_datetime(dados["start_date"]).dt.to_period("M")
    duplicados = dados.duplicated(["hydro_id", "periodo"], keep=False)
    if duplicados.any():
        exemplo = dados.loc[duplicados].iloc[0]
        raise ValueError(
            f"O histórico possui mais de um registro para hydro_id {int(exemplo['hydro_id'])} "
            f"no mês {exemplo['periodo']}."
        )
    resultados = {}
    for hydro_id, grupo in dados.groupby("hydro_id", sort=True):
        serie = grupo.set_index("periodo")["vazao"].sort_index()
        indice_completo = pd.period_range(serie.index.min(), serie.index.max(), freq="M")
        serie = pd.to_numeric(serie.reindex(indice_completo), errors="coerce")
        serie = serie.where(np.isfinite(serie))
        conjuntos = {
            JANELAS[0]: serie,
            JANELAS[1]: serie.rolling(6, min_periods=6).sum(),
            JANELAS[2]: serie.rolling(12, min_periods=12).sum(),
        }
        for janela, valores in conjuntos.items():
            for mes_num in range(1, 13):
                selecao = valores.loc[valores.index.month == mes_num]
                resultados[(int(hydro_id), janela, mes_num)] = _resumir_valores(selecao)
    return resultados


def gerar_relatorio(
    caminho_cenarios: str,
    caminho_historico: str,
    nomes_por_id: dict[int, str],
    caminho_stages: str,
) -> pd.DataFrame:
    """Entrega os resumos sem soma e das janelas de 6 e 12 para todas as UHEs."""
    calendario = _carregar_calendario(caminho_stages)
    cenarios = _resumo_cenarios(caminho_cenarios, calendario)
    historico = _resumo_historico(caminho_historico)
    por_usina = {int(chave): grupo for chave, grupo in cenarios.groupby("hydro_id", sort=False)}
    ids = sorted(set(por_usina) | {hydro_id for hydro_id, _, _ in historico})
    ordem_estagios = calendario.set_index("stage_id")["ordem"].to_dict()
    linhas = []
    for hydro_id in ids:
        nome = nomes_por_id.get(hydro_id, f"Usina sem nome — hydro_id {hydro_id}")
        for janela in JANELAS:
            for mes_num in range(1, 13):
                resumo = historico.get((hydro_id, janela, mes_num), {"n_validos": 0})
                linhas.append({
                    "usina": nome, "hydro_id": hydro_id, "fonte": "Histórico",
                    "janela": janela, "mes_num": mes_num,
                    "mes_nome": MESES[mes_num], "stage_id": None,
                    **{chave: resumo.get(chave) for chave in (
                        "n_validos", "media", "desvio", "assimetria", "curtose", "minimo", "maximo"
                    )},
                })
        if hydro_id not in por_usina:
            continue
        dados_usina = por_usina[hydro_id]
        for janela in JANELAS:
            grupo_janela = dados_usina.loc[dados_usina["janela"].eq(janela)].copy()
            grupo_janela["_ordem"] = grupo_janela["stage_id"].map(ordem_estagios)
            grupo_janela = grupo_janela.sort_values("_ordem")
            for registro in grupo_janela.itertuples(index=False):
                desvio = _numero_finito(registro.desvio)
                n = int(registro.n_validos)
                m3 = _numero_finito(registro.m3)
                m4 = _numero_finito(registro.m4)
                assimetria = m3 / desvio ** 3 if n >= 2 and desvio and m3 is not None else None
                curtose = m4 / desvio ** 4 - 3 if n >= 2 and desvio and m4 is not None else None
                linhas.append({
                    "usina": nome, "hydro_id": hydro_id, "fonte": "Cenários",
                    "janela": janela, "mes_num": None, "mes_nome": "",
                    "stage_id": int(registro.stage_id),
                    "n_validos": n,
                    "media": _numero_finito(registro.media),
                    "desvio": desvio,
                    "assimetria": _numero_finito(assimetria),
                    "curtose": _numero_finito(curtose),
                    "minimo": _numero_finito(registro.minimo),
                    "maximo": _numero_finito(registro.maximo),
                })
    return pd.DataFrame(linhas)


def filtrar_relatorio(
    relatorio: pd.DataFrame,
    usinas: list[int] | None = None,
    fonte: str = "Todas",
    meses: list[int] | None = None,
    estagios: list[int] | None = None,
    janelas: list[str] | None = None,
) -> pd.DataFrame:
    """Filtra cada tipo de linha pelo seu próprio calendário e preserva a ordem."""
    resultado = relatorio
    if usinas:
        resultado = resultado.loc[resultado["hydro_id"].isin(usinas)]
    if fonte != "Todas":
        resultado = resultado.loc[resultado["fonte"].eq(fonte)]
    if janelas:
        resultado = resultado.loc[resultado["janela"].isin(janelas)]
    if meses:
        resultado = resultado.loc[
            resultado["fonte"].ne("Histórico") | resultado["mes_num"].isin(meses)
        ]
    if estagios:
        resultado = resultado.loc[
            resultado["fonte"].ne("Cenários") | resultado["stage_id"].isin(estagios)
        ]
    return resultado.reset_index(drop=True)


def tabela_visivel(relatorio: pd.DataFrame) -> pd.DataFrame:
    """Aplica nomes legíveis sem arredondar os valores usados na exportação."""
    visivel = relatorio[[chave for chave, _ in COLUNAS]].copy()
    visivel.columns = [nome for _, nome in COLUNAS]
    visivel["stage_id"] = visivel["stage_id"].astype("Int64")
    visivel["Observações válidas"] = visivel["Observações válidas"].astype("Int64")
    return visivel


def _decimal_visivel(valor) -> str | None:
    """Garante ponto decimal e até duas casas no relatório visível do Excel."""
    numero = _numero_finito(valor)
    if numero is None:
        return None
    texto = f"{numero:.2f}".rstrip("0").rstrip(".")
    return "0" if texto in ("-0", "") else texto


def exportar_xlsx(relatorio: pd.DataFrame) -> bytes:
    """Cria o relatório ciano e conserva números integrais em uma aba auxiliar."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    planilha = livro.active
    planilha.title = "Resumo estatístico"
    numerica = livro.create_sheet("Valores numéricos")
    cabecalhos = [nome for _, nome in COLUNAS]
    for folha in (planilha, numerica):
        folha.append(cabecalhos)
        folha.sheet_view.showGridLines = False
        folha.freeze_panes = "C2"

    ciano = "09B6CB"
    ciano_escuro = "067F94"
    ciano_claro = "ECF9FB"
    borda_fina = Side(style="thin", color="B3DEE5")
    borda_divisoria = Side(style="medium", color=ciano_escuro)
    borda_padrao = Border(left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina)
    borda_separacao = Border(left=borda_fina, right=borda_fina, top=borda_divisoria, bottom=borda_fina)
    alinhamento_texto = Alignment(horizontal="left", vertical="center")
    alinhamento_numero = Alignment(horizontal="right", vertical="center")
    fonte_cabecalho = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
    fonte_corpo = Font(name="Aptos", size=10, color="183444")
    preenchimento_cabecalho = PatternFill("solid", fgColor=ciano)
    preenchimento_historico = PatternFill("solid", fgColor=ciano_claro)
    preenchimento_cenarios = PatternFill("solid", fgColor="FFFFFF")
    larguras = [len(nome) + 2 for nome in cabecalhos]
    chaves_numericas_excel = {"hydro_id", "stage_id", "n_validos"} | COLUNAS_NUMERICAS
    posicoes_numericas = {
        indice for indice, (chave, _) in enumerate(COLUNAS, start=1)
        if chave in chaves_numericas_excel
    }

    for celula in planilha[1]:
        celula.fill = preenchimento_cabecalho
        celula.font = fonte_cabecalho
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celula.border = borda_padrao
    planilha.row_dimensions[1].height = 32

    anterior = None
    for indice, registro in enumerate(relatorio.itertuples(index=False), start=2):
        dados = registro._asdict()
        valores_numericos = [dados[chave] for chave, _ in COLUNAS]
        numerica.append([
            None if valor is None or (isinstance(valor, float) and not math.isfinite(valor)) else valor
            for valor in valores_numericos
        ])
        valores_visiveis = [
            _decimal_visivel(dados[chave]) if chave in COLUNAS_NUMERICAS else dados[chave]
            for chave, _ in COLUNAS
        ]
        planilha.append(valores_visiveis)
        divisoria = (
            anterior is not None
            and anterior[0] == dados["hydro_id"]
            and anterior[1] == "Histórico"
            and dados["fonte"] == "Cenários"
        )
        preenchimento = preenchimento_historico if dados["fonte"] == "Histórico" else preenchimento_cenarios
        for posicao in range(1, len(COLUNAS) + 1):
            celula = planilha.cell(row=indice, column=posicao)
            celula.font = fonte_corpo
            celula.fill = preenchimento
            celula.alignment = alinhamento_numero if posicao in posicoes_numericas else alinhamento_texto
            celula.border = borda_separacao if divisoria else borda_padrao
            larguras[posicao - 1] = max(larguras[posicao - 1], len(str(celula.value or "")) + 2)
        planilha.row_dimensions[indice].height = 19
        anterior = (dados["hydro_id"], dados["fonte"])

    for posicao, largura in enumerate(larguras, start=1):
        planilha.column_dimensions[get_column_letter(posicao)].width = min(max(largura, 12), 80)
    ultima_coluna = get_column_letter(len(COLUNAS))
    planilha.auto_filter.ref = f"A1:{ultima_coluna}{max(1, planilha.max_row)}"
    planilha.sheet_properties.pageSetUpPr.fitToPage = True
    planilha.sheet_properties.tabColor = ciano_escuro
    planilha.print_options.horizontalCentered = True
    numerica.sheet_state = "hidden"  # Preserva números reais para operações sem mudar o ponto decimal do relatório.

    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def nome_arquivo_exportacao() -> str:
    return f"relatorio_estatistico_cobre_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
