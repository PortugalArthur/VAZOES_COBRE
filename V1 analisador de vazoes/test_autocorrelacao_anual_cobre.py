"""Casos controlados da soma anual e da autocorrelação de lag 1 e 2."""

from datetime import date
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

import duckdb
import numpy as np
import pandas as pd
import pytest

from preparar_ambiente import garantir_dependencias

garantir_dependencias({"openpyxl": "openpyxl>=3.1"})
from openpyxl import load_workbook

from autocorrelacao_anual_cobre import (
    exportar_autocorrelacao_anual_xlsx,
    filtrar_autocorrelacao_anual,
    gerar_autocorrelacao_anual,
)


def _fixture(pasta: Path, inicio: date = date(2025, 8, 1),
             fim: date = date(2029, 3, 1), duplicar: bool = False):
    meses = pd.date_range(inicio, fim, freq="MS")
    calendario = {
        100 + indice * 3: {"inicio": instante.date(), "ordem": indice}
        for indice, instante in enumerate(meses)
    }
    totais = {
        2026: [1.0, 2.0, 3.0, 4.0],
        2027: [4.0, 3.0, 2.0, 1.0],
        2028: [1.0, -1.0, -1.0, 1.0],
    }
    linhas = []
    for stage_id, dados in calendario.items():
        ano, mes = dados["inicio"].year, dados["inicio"].month
        for hydro_id in (10, 20):
            for scenario_id in range(1, 5):
                if hydro_id == 10 and ano in totais:
                    valor = totais[ano][scenario_id - 1] / 12
                else:
                    valor = 5.0
                if hydro_id == 10 and ano == 2027 and mes == 6 and scenario_id == 4:
                    valor = np.nan
                linhas.append({
                    "hydro_id": hydro_id, "scenario_id": scenario_id,
                    "stage_id": stage_id,
                    "node_id": stage_id * 100 + hydro_id + scenario_id,
                    "incremental_inflow_m3s": valor,
                })
    dados = pd.DataFrame(linhas).sample(frac=1, random_state=7).reset_index(drop=True)
    if duplicar:
        dados = pd.concat([dados, dados.iloc[[0]]], ignore_index=True)
    caminho = pasta / "anual.parquet"
    conexao = duckdb.connect()
    try:
        conexao.register("fixture", dados)
        conexao.execute("COPY fixture TO ? (FORMAT PARQUET)", [str(caminho)])
    finally:
        conexao.close()
    return caminho, calendario


def test_anos_civis_pares_faltantes_tabela_e_xlsx():
    with TemporaryDirectory() as temporario:
        caminho, calendario = _fixture(Path(temporario))
        tabela, detalhe = gerar_autocorrelacao_anual(
            str(caminho), {10: "Teste", 20: "Constante"}, calendario
        )
        assert tabela.columns.tolist() == ["UHE", "hydro_id", "Ano", "Lag 1", "Lag 2"]
        assert tabela.shape == (6, 5)
        assert tabela.loc[tabela.hydro_id.eq(10), "Ano"].tolist() == [2026, 2027, 2028]
        assert pd.isna(tabela.loc[tabela.hydro_id.eq(10) & tabela.Ano.eq(2026), "Lag 1"]).all()
        assert tabela.loc[tabela.hydro_id.eq(10) & tabela.Ano.eq(2027), "Lag 1"].iloc[0] == pytest.approx(-1)
        assert tabela.loc[tabela.hydro_id.eq(10) & tabela.Ano.eq(2028), "Lag 2"].iloc[0] == pytest.approx(0, abs=1e-14)
        par1 = detalhe.loc[detalhe.hydro_id.eq(10) & detalhe.Ano.eq(2027) & detalhe.Lag.eq(1)].iloc[0]
        assert par1["Pares válidos"] == 3
        assert par1["Somas completas atuais"] == 3
        assert par1["Média atual (soma m³/s)"] == pytest.approx(3)
        assert par1["Média passada (soma m³/s)"] == pytest.approx(2)
        par2 = detalhe.loc[detalhe.hydro_id.eq(10) & detalhe.Ano.eq(2028) & detalhe.Lag.eq(2)].iloc[0]
        assert par2["Pares válidos"] == 4
        assert par2["Status"] == "calculado"
        assert detalhe.loc[detalhe.hydro_id.eq(20) & detalhe.Ano.eq(2027) & detalhe.Lag.eq(1), "Motivo"].iloc[0].find("constante") >= 0
        filtrada, auditada = filtrar_autocorrelacao_anual(tabela, detalhe, [10], [2028])
        assert filtrada.shape == (1, 5)
        assert len(auditada) == 2
        livro = load_workbook(BytesIO(exportar_autocorrelacao_anual_xlsx(filtrada, auditada)))
        assert livro["Autocorrelação anual"].max_row == 2
        assert livro["Valores completos"].sheet_state == "hidden"
        assert livro["Diagnóstico"].max_row == 3
        assert livro["Autocorrelação anual"]["E2"].value == "0.00"
        assert livro["Autocorrelação anual"]["D2"].value == "0.87"
        assert livro["Valores completos"]["D2"].value == pytest.approx(
            filtrada["Lag 1"].iloc[0]
        )


def test_inicio_em_janeiro_e_erro_de_chave_duplicada():
    with TemporaryDirectory() as temporario:
        pasta = Path(temporario)
        caminho, calendario = _fixture(
            pasta, date(2026, 1, 1), date(2028, 12, 1)
        )
        tabela, _ = gerar_autocorrelacao_anual(str(caminho), {10: "Teste"}, calendario, 10)
        assert tabela["Ano"].tolist() == [2026, 2027, 2028]
        assert pd.isna(tabela.loc[tabela.Ano.eq(2026), "Lag 1"]).all()
        assert tabela.loc[tabela.Ano.eq(2027), "Lag 1"].iloc[0] == pytest.approx(-1)
        caminho_duplicado, calendario = _fixture(
            pasta, date(2026, 1, 1), date(2028, 12, 1), duplicar=True
        )
        with pytest.raises(ValueError, match="duplicada"):
            gerar_autocorrelacao_anual(str(caminho_duplicado), {}, calendario)
